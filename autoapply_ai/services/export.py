"""
autoapply/services/export.py
Generate PDF and DOCX bytes from text content.
"""
from __future__ import annotations
from io import BytesIO


# ── DOCX ──────────────────────────────────────────────────────

def export_docx(title: str, content: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Style heading
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)

    doc.add_paragraph("")

    for line in content.split("\n"):
        p = doc.add_paragraph(line if line.strip() else "")
        p.paragraph_format.space_after = Pt(2)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ── PDF ───────────────────────────────────────────────────────

def export_pdf(title: str, content: str) -> bytes:
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=20*mm, rightMargin=20*mm,
        topMargin=20*mm, bottomMargin=20*mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        fontSize=16, spaceAfter=12,
        textColor=colors.HexColor("#1A1A1A"),
    )
    body_style = ParagraphStyle(
        "CustomBody",
        parent=styles["BodyText"],
        fontSize=10, leading=14,
        textColor=colors.HexColor("#333333"),
        spaceAfter=4,
    )

    story = [Paragraph(title, title_style), Spacer(1, 8*mm)]
    for line in content.split("\n"):
        story.append(Paragraph(line.strip() or "&nbsp;", body_style))

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


# ── Base64 helpers (for data-URI downloads) ───────────────────

def to_data_uri_pdf(title: str, content: str) -> str:
    import base64
    data = export_pdf(title, content)
    b64  = base64.b64encode(data).decode()
    return f"data:application/pdf;base64,{b64}"


def to_data_uri_docx(title: str, content: str) -> str:
    import base64
    data = export_docx(title, content)
    b64  = base64.b64encode(data).decode()
    return "data:application/vnd.openxmlformats-officedocument.wordprocessingml.document;base64," + b64